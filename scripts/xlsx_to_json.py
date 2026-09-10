#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel(.xlsx) 소장 게임 목록 -> js/data.js 변환 스크립트

사용법:
  python3 xlsx_to_json.py 리스트.xlsx

전제:
  - '리스트' 시트의 B58 행부터 시작하는 표(순번 / 타이틀 / 구분 / 형태 / 엔딩 / 비고)를 읽습니다.
  - 새 게임이 목록 중간에 추가되어도, 카테고리 + 정규화된 제목을 기준으로
    기존과 동일한 id가 생성되는 게임은 id가 유지됩니다.
    (완전히 새로운 게임에는 새 id가 발급됩니다.)
  - 이 스크립트를 실행하면 js/data.js 파일이 갱신됩니다.
    클리어/즐겨찾기 기록은 브라우저 localStorage에 id 기준으로 저장되므로
    번호가 밀리거나 목록 순서가 바뀌어도 기록은 그대로 유지됩니다.

필요 패키지:
  pip install openpyxl
"""

import sys
import json
import re
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl이 필요합니다: pip install openpyxl")
    sys.exit(1)

# 구분 코드 -> (내부 key, 화면 표시 라벨, 색상)
CAT_META = {
    'S2':  ('s2',        'Switch2',          '#782a2a'),
    'S2D': ('s2d',       'Switch2 DL',       '#78492a'),
    1:     ('g1st',      '1st',              '#78692a'),
    2:     ('g2nd',      '2nd',              '#69782a'),
    3:     ('g3rd',      '3rd',              '#49782a'),
    4:     ('g4th',      '4th',              '#2a782a'),
    5:     ('g5th',      '5th',              '#2a7849'),
    6:     ('g6th',      '6th',              '#2a7869'),
    7:     ('g7th',      '7th',              '#2a6978'),
    8:     ('action',    'Action&Arcade',    '#2a4978'),
    9:     ('indie',     'Indie Action',     '#2a2a78'),
    10:    ('rpg',       'RPG&Simulation',   '#492a78'),
    11:    ('adventure', 'Adventure&Puzzle', '#692a78'),
    12:    ('minor',     'Minor&Etc.',       '#782a69'),
    13:    ('multi',     'Multi&Repeat',     '#782a49'),
}

CATEGORY_ORDER = ['S2','S2D',1,2,3,4,5,6,7,8,9,10,11,12,13]


def slugify(title: str) -> str:
    if not title:
        return "untitled"
    s = re.sub(r'[^0-9A-Za-z가-힣]+', '-', title).strip('-').lower()
    return s or "untitled"


def extract_games(xlsx_path: str):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb['리스트']

    items = []
    for r in range(59, ws.max_row + 1):
        no = ws.cell(row=r, column=2).value
        title = ws.cell(row=r, column=3).value
        cat = ws.cell(row=r, column=4).value
        ending = ws.cell(row=r, column=6).value  # 엔딩(클리어) 컬럼
        if title is None and cat is None:
            continue
        if cat not in CAT_META:
            continue
        items.append({
            "no": no,
            "title": title.strip() if isinstance(title, str) else title,
            "cat": cat,
            "done": (ending == 'Y'),
        })

    seen = {}
    games = []
    for it in items:
        key, label, color = CAT_META[it['cat']]
        base_id = f"{key}-{slugify(it['title'])}"
        n = seen.get(base_id, 0) + 1
        seen[base_id] = n
        gid = base_id if n == 1 else f"{base_id}-{n}"
        games.append({
            "id": gid,
            "number": it['no'],
            "title": it['title'],
            "category": key,
            "d0": it['done'],
        })
    return games


def build_categories():
    return [
        {"key": key, "label": label, "color": color}
        for cat_code in CATEGORY_ORDER
        for (key, label, color) in [CAT_META[cat_code]]
    ]


def main():
    if len(sys.argv) < 2:
        print("사용법: python3 xlsx_to_json.py <엑셀파일.xlsx>")
        sys.exit(1)

    xlsx_path = sys.argv[1]
    games = extract_games(xlsx_path)
    categories = build_categories()

    out_path = Path(__file__).resolve().parent.parent / "js" / "data.js"
    js = (
        "// 자동 생성 파일 - Excel에서 추출한 게임 데이터\n"
        "// 목록을 업데이트하려면 이 파일만 교체하면 됩니다.\n"
        "// (scripts/xlsx_to_json.py 참고)\n\n"
        f"const CATEGORIES = {json.dumps(categories, ensure_ascii=False, separators=(',',':'))};\n\n"
        f"const GAMES = {json.dumps(games, ensure_ascii=False, separators=(',',':'))};\n"
    )
    out_path.write_text(js, encoding="utf-8")
    print(f"완료: {len(games)}개 게임을 {out_path} 에 저장했습니다.")


if __name__ == "__main__":
    main()
